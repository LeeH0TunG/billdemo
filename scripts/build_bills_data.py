"""Convert the supplied workbook into browser data for this static demo."""
import json
from collections import defaultdict
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
workbook = openpyxl.load_workbook(ROOT / "assets" / "demo_bills.xlsx", data_only=True)
sheet = workbook["账单明细"]
headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
cards = {str(key): value for key, value in workbook["Sheet1"].iter_rows(min_row=2, values_only=True) if key and value}
months = defaultdict(list)
for row in sheet.iter_rows(min_row=2, values_only=True):
    item = dict(zip(headers, row))
    if item["账单月份"] and item["人民币金额"] is not None: months[item["账单月份"]].append(item)
def money(value): return f"¥ {'-' if value < 0 else ''}{abs(value):,.2f}"
ordered = sorted(months, reverse=True); totals = {month: sum(item["人民币金额"] for item in rows if item["交易类型"] != "还款") for month, rows in months.items()}; transactions = {}; bills = []; last_year = None
latest = next(month for month in ordered if month != "预估未出账单")
for index, month in enumerate(ordered):
    pending = month == "预估未出账单"
    year = month[:4]
    if not pending and year != last_year: bills.append({"year": year + "年"}); last_year = year
    key = "pending" if pending else "current" if month == latest else f"{int(month[5:])}月"
    rows = sorted(months[month], key=lambda item: item["交易日"], reverse=True); records = []; last_day = None
    for item in rows:
        day = item["交易日"]
        if day != last_day: records.append({"date": f"{int(day[5:7])}月{int(day[8:10])}日"}); last_day = day
        amount = item["人民币金额"]; merchant = item["交易摘要"]
        records.append({"merchant": merchant, "tail": f"尾号({item['卡号末四位']})", "amount": money(amount), "type": item["交易类型"], "time": f"{day} {item['交易时间'].isoformat(timespec='seconds')}" if item["交易时间"] else day, "company": item["公司名"], "posted": item["记账日"], "channel": merchant.split("-")[0] if "-" in merchant else "无", "category": item["消费分类"], "card": item["完整卡号"] or cards.get(str(item["卡号末四位"])), "country": item["交易国家"] or "中国", "originalAmount": item["交易地金额"], "originalCurrency": item["交易地币种"], "refund": item["交易类型"] == "退款"})
    if pending:
        bills.append({"id": key, "title": month, "amount": money(totals[month]), "info": True, "tx": records})
        transactions[key] = records
        continue
    total = totals[month]; number = int(month[5:]); prior = totals.get(ordered[index + 1]) if index + 1 < len(ordered) else None
    bill = {"id": key, "title": "本期账单" if month == latest else f"{number}月账单", "amount": money(total), "date": f"{number - 1 or 12:02d}/15-{number:02d}/14", "tx": records}
    if prior not in (None, 0):
        change = (total - prior) / abs(prior); direction = "up" if change >= 0 else "down"; bill["meta"] = f'比上月 <b class="{direction}">{"↑" if change >= 0 else "↓"} {abs(change):.0%}</b>'
    bills.append(bill); transactions[key] = records
(ROOT / "assets" / "bills-data.js").write_text("window.billDemoData=" + json.dumps({"transactions": transactions, "bills": bills}, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
